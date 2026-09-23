# -*- coding: utf-8 -*-
"""
AI Mastery Workshop — deck generator (MakeFirstMillion series)
Builds a 16:9 .pptx from python-pptx implementing the brand design system:
  - Primary accent:  #FF7D00 (orange)
  - Background base: #FFF9F4 (warm off-white) on most slides
  - Secondary/wins:  #62C41F (green)
  - Near-black text: #1A1A1A
  - High-contrast "hook" slides: near-black bg + orange/off-white text
Run:  py build/generate_deck.py
Output: output/AI-Mastery-Workshop.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette --
ORANGE = RGBColor(0xFF, 0x7D, 0x00)
BG_LIGHT = RGBColor(0xFF, 0xF9, 0xF4)
GREEN = RGBColor(0x62, 0xC4, 0x1F)
INK = RGBColor(0x1A, 0x1A, 0x1A)
DARK_BG = RGBColor(0x14, 0x13, 0x12)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x64, 0x5C)          # muted body text on light bg
MUTED_ON_DARK = RGBColor(0xC9, 0xC2, 0xB8)
CARD_BORDER = RGBColor(0xEA, 0xDE, 0xCF)
CODE_BG = RGBColor(0x1E, 0x1B, 0x18)
CODE_GREEN = RGBColor(0x8FE39A)  if False else RGBColor(0x8F, 0xE3, 0x9A)
CODE_ORANGE = RGBColor(0xFF, 0xB0, 0x5C)
CODE_TEXT = RGBColor(0xE8, 0xE3, 0xDC)

HEAD_FONT = "Montserrat"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_SECTIONS = 13

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]

# ------------------------------------------------------------- utilities --

def add_slide(bg=BG_LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    # send to back
    sp = r._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def _set_no_autofit(tf):
    el = tf._txBody
    bodyPr = el.find(qn('a:bodyPr'))
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        e = bodyPr.find(qn(tag))
        if e is not None:
            bodyPr.remove(e)
    bodyPr.append(el.makeelement(qn('a:noAutofit'), {}))


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False,
            font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            spacing=None, line_spacing=1.0, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    _set_no_autofit(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color
        if spacing is not None:
            _letter_spacing(run, spacing)
    return tb


def _letter_spacing(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(int(pts * 100)))


def rich_textbox(slide, x, y, w, h, runs, size=18, font=BODY_FONT,
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    """runs: list of (text, color, bold) rendered on one paragraph/line."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    _set_no_autofit(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for text, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=Pt(1), shadow=False, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = line_w
    sh.shadow.inherit = False
    if shadow:
        _soft_shadow(sh)
    return sh


def _soft_shadow(shape):
    spPr = shape._element.spPr
    existing = spPr.find(qn('a:effectLst'))
    if existing is not None:
        spPr.remove(existing)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:outerShdw blurRad="180000" dist="30000" dir="5400000" rotWithShape="0">'
        '<a:srgbClr val="1A1A1A"><a:alpha val="18000"/></a:srgbClr>'
        '</a:outerShdw></a:effectLst>'
    )
    from lxml import etree
    spPr.append(etree.fromstring(xml))


def kicker(slide, text, x=Inches(0.7), y=Inches(0.5), color=ORANGE, on_dark=False):
    textbox(slide, x, y, Inches(9), Inches(0.4), text.upper(), size=13, color=color,
            bold=True, font=BODY_FONT, spacing=2.2)


def progress_tag(slide, section_num, total=TOTAL_SECTIONS, on_dark=False):
    c = MUTED_ON_DARK if on_dark else MUTED
    textbox(slide, Inches(10.6), Inches(0.5), Inches(2.2), Inches(0.35),
            f"SECTION {section_num} OF {total}", size=11, color=c, bold=True,
            align=PP_ALIGN.RIGHT, spacing=1.4)


def underline(slide, x, y, w, color=ORANGE, h=Pt(6)):
    rect(slide, x, y, w, h, fill=color)


def headline(slide, text, x=Inches(0.7), y=Inches(1.1), w=Inches(11.9), size=44,
             color=INK, align=PP_ALIGN.LEFT):
    return textbox(slide, x, y, w, Inches(2.2), text, size=size, color=color,
                    bold=True, font=HEAD_FONT, align=align, line_spacing=1.02)


def live_badge(slide, on_dark=True):
    y = Inches(6.9)
    dot_color = ORANGE
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), y + Inches(0.06), Inches(0.12), Inches(0.12))
    dot.fill.solid(); dot.fill.fore_color.rgb = dot_color
    dot.line.fill.background(); dot.shadow.inherit = False
    color = BG_LIGHT if on_dark else INK
    textbox(slide, Inches(0.92), y, Inches(7), Inches(0.35),
            "LIVE SESSION  |  MAKEFIRSTMILLION SERIES", size=11, color=color,
            bold=True, spacing=1.4)


def footer_page(slide, n):
    textbox(slide, Inches(12.5), Inches(7.08), Inches(0.6), Inches(0.3), str(n),
            size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def stat_card(slide, x, y, w, h, value, label, value_color=ORANGE):
    rect(slide, x, y, w, h, fill=WHITE, line_color=CARD_BORDER, line_w=Pt(1),
         shadow=True, radius=0.06)
    textbox(slide, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.7),
            value, size=26, color=value_color, bold=True, font=HEAD_FONT)
    textbox(slide, x + Inches(0.25), y + h - Inches(0.65), w - Inches(0.5), Inches(0.5),
            label, size=13, color=MUTED, bold=False)


def code_block(slide, x, y, w, h, lines, label="PROMPT", title_color=CODE_ORANGE):
    rect(slide, x, y, w, h, fill=CODE_BG, radius=0.03, shadow=True)
    # traffic-light dots
    for i, c in enumerate([RGBColor(0xFF, 0x5F, 0x56), RGBColor(0xFF, 0xBD, 0x2E), RGBColor(0x27, 0xC9, 0x3F)]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22) + Inches(0.22) * i, y + Inches(0.18), Inches(0.12), Inches(0.12))
        dot.fill.solid(); dot.fill.fore_color.rgb = c
        dot.line.fill.background(); dot.shadow.inherit = False
    textbox(slide, x + w - Inches(2.3), y + Inches(0.14), Inches(2.05), Inches(0.3),
            label.upper(), size=10.5, color=title_color, bold=True, font=MONO_FONT,
            align=PP_ALIGN.RIGHT, spacing=1.2)
    body = "\n".join(lines)
    textbox(slide, x + Inches(0.3), y + Inches(0.55), w - Inches(0.6), h - Inches(0.75),
            body, size=13, color=CODE_TEXT, font=MONO_FONT, line_spacing=1.25)


def bullets(slide, items, x, y, w, h, size=16, color=INK, gap=0.12, marker_color=ORANGE, bold_marker=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    _set_no_autofit(tf)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap * 72)
        p.line_spacing = 1.15
        r1 = p.add_run(); r1.text = "—  "
        r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = marker_color; r1.font.name = BODY_FONT
        r2 = p.add_run(); r2.text = item
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = BODY_FONT
    return tb


def placeholder_box(slide, x, y, w, h, label):
    rect(slide, x, y, w, h, fill=RGBColor(0xF0, 0xE8, 0xDC), line_color=ORANGE, line_w=Pt(1.25), radius=0.05)
    textbox(slide, x, y + h/2 - Inches(0.25), w, Inches(0.5), f"[ {label} ]", size=13,
            color=ORANGE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def section_break(slide_bg, section_num, kicker_text, title, subtitle=None):
    s = add_slide(bg=DARK_BG)
    kicker(s, kicker_text, color=ORANGE)
    progress_tag(s, section_num, on_dark=True)
    headline(s, title, y=Inches(2.7), size=48, color=BG_LIGHT)
    underline(s, Inches(0.72), Inches(4.35), Inches(2.1), color=ORANGE)
    if subtitle:
        textbox(s, Inches(0.7), Inches(4.65), Inches(10.5), Inches(1.2), subtitle,
                size=18, color=MUTED_ON_DARK, line_spacing=1.3)
    live_badge(s, on_dark=True)
    return s


def content_header(s, section_num, kicker_text, title, title_size=34):
    kicker(s, kicker_text)
    progress_tag(s, section_num)
    headline(s, title, y=Inches(0.95), size=title_size)
    underline(s, Inches(0.72), Inches(0.95) + Inches(0.02) + Emu(int(title_size * 12700 * 1.15)), Inches(1.6))


page_num = [0]

def new_content_slide(section_num, kicker_text, title, title_size=32, bg=BG_LIGHT):
    s = add_slide(bg=bg)
    kicker(s, kicker_text)
    progress_tag(s, section_num)
    headline(s, title, y=Inches(0.95), size=title_size)
    page_num[0] += 1
    footer_page(s, page_num[0])
    return s


def step_list(slide, steps, x, y, w, row_h=0.92, circle_d=0.5, size=16, color=INK, start=1):
    """Numbered-circle step list — same visual pattern as the job-hunting /
    curriculum steps in the legacy deck."""
    for i, st in enumerate(steps):
        yy = y + Inches(row_h) * i
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, yy, Inches(circle_d), Inches(circle_d))
        circ.fill.solid(); circ.fill.fore_color.rgb = ORANGE
        circ.line.fill.background(); circ.shadow.inherit = False
        tf = circ.text_frame; tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i + start)
        r.font.bold = True; r.font.size = Pt(16); r.font.color.rgb = WHITE; r.font.name = HEAD_FONT
        textbox(slide, x + Inches(circle_d + 0.28), yy, w - Inches(circle_d + 0.28), Inches(circle_d),
                st, size=size, color=color, anchor=MSO_ANCHOR.MIDDLE)


def chip_grid(slide, items, x, y, w, h, cols=3, size=13, color=INK, fill=WHITE, border=CARD_BORDER):
    """Grid of short text chips — mirrors the deck's `.uc-grid` pattern."""
    n = len(items)
    rows = -(-n // cols)
    gap = Inches(0.2)
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, item in enumerate(items):
        col, row = i % cols, i // cols
        xx = x + col * (cw + gap)
        yy = y + row * (ch + gap)
        rect(slide, xx, yy, cw, ch, fill=fill, line_color=border, line_w=Pt(1), radius=0.12)
        textbox(slide, xx + Inches(0.18), yy, cw - Inches(0.36), ch, item, size=size, color=color,
                anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.12)


def qa_row(slide, x, y, w, label, tag_text, why_text, power=False, h=1.0):
    """Weak-vs-power comparison row — mirrors the deck's `.qa-item` pattern."""
    fill = RGBColor(0xEE, 0xF8, 0xE4) if power else WHITE
    border = GREEN if power else CARD_BORDER
    tag_color = GREEN if power else RGBColor(0xC4, 0x4A, 0x2A)
    rect(slide, x, y, w, Inches(h), fill=fill, line_color=border, line_w=Pt(1.25), radius=0.08, shadow=True)
    textbox(slide, x + Inches(0.3), y + Inches(0.12), w - Inches(0.6), Inches(0.4), label,
            size=16, color=INK, bold=True)
    rich_textbox(slide, x + Inches(0.3), y + Inches(0.55), w - Inches(0.6), Inches(0.4),
                 [(tag_text + "  —  ", tag_color, True), (why_text, MUTED, False)], size=13)


def table_block(slide, x, y, w, h, headers, rows):
    """Simple comparison table (native pptx table)."""
    n_rows = len(rows) + 1
    tbl_shape = slide.shapes.add_table(n_rows, len(headers), x, y, w, h)
    tbl = tbl_shape.table
    for j, htext in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = htext
        cell.fill.solid(); cell.fill.fore_color.rgb = DARK_BG
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.size = Pt(15); r.font.color.rgb = ORANGE; r.font.name = HEAD_FONT
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xF6, 0xF0, 0xE8)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(14); r.font.color.rgb = INK; r.font.name = BODY_FONT
    return tbl_shape


# =============================================================================
# TITLE — HOOK & HOUSEKEEPING
# =============================================================================

s = add_slide(bg=DARK_BG)
kicker(s, "AI MASTERY WORKSHOP", color=ORANGE)
headline(s, "AI IS NOT\nOPTIONAL\nANYMORE", y=Inches(1.5), size=64, color=BG_LIGHT)
underline(s, Inches(0.72), Inches(4.55), Inches(2.6), color=ORANGE, h=Pt(8))
textbox(s, Inches(0.72), Inches(4.85), Inches(9.5), Inches(1.3),
        "You will leave knowing how to use AI, build with AI, automate with AI,\n"
        "and start thinking about how to make money with AI.",
        size=18, color=MUTED_ON_DARK, line_spacing=1.3)
rect(s, Inches(0.72), Inches(5.9), Inches(8.6), Inches(0.9), fill=RGBColor(0x24, 0x20, 0x1C), radius=0.12)
textbox(s, Inches(1.0), Inches(6.02), Inches(8.1), Inches(0.7),
        "“Awaz clear aa rahi hai?”\n(Confirming audio is clear — drop a comment.)",
        size=14, color=ORANGE, bold=False, line_spacing=1.2)
live_badge(s, on_dark=True)
footer_page(s, 1); page_num[0] = 1

s = new_content_slide(1, "By the End of Today", "You Will Be Able To", title_size=34)
byend = ["AI, ML, Deep Learning, GenAI & AGI — explained simply", "A repeatable prompting framework",
         "Research with an agentic browser", "Your own custom AI assistant / GPT",
         "A working app, built with Claude", "Agent vs. automation — the real difference",
         "AI videos & avatars with HeyGen", "A real automation workflow in n8n",
         "Turning these skills into income"]
chip_grid(s, byend, Inches(0.72), Inches(2.0), Inches(11.9), Inches(4.5), cols=3, size=14)
textbox(s, Inches(0.72), Inches(6.75), Inches(11.9), Inches(0.5),
        "This is not a “100 AI tools” tour — the goal is to understand how to think about AI and turn tools into outcomes.",
        size=13, color=MUTED, bold=True)

# =============================================================================
# SECTION 2 — CREDIBILITY
# =============================================================================

s = new_content_slide(2, "Credibility", "Why Listen to Me?", title_size=36)
placeholder_box(s, Inches(0.72), Inches(2.0), Inches(3.1), Inches(3.6), "PRESENTER PHOTO")
textbox(s, Inches(4.1), Inches(2.15), Inches(4.5), Inches(0.6), "AI Engineer / Practitioner", size=22, color=ORANGE, bold=True, font=HEAD_FONT)
textbox(s, Inches(4.1), Inches(2.75), Inches(8.3), Inches(1.3),
        "Runs an AI startup studio building products and automations for US clients.\n"
        "Built Aaghaz as an AI education and startup platform focused on\n"
        "practical, hands-on skills.",
        size=14, color=MUTED, line_spacing=1.3)
cards = [("15,000+", "Students trained"), ("5+ yrs", "In the AI field"),
         ("AI Startup", "Studio for US clients"), ("Aaghaz", "AI education platform")]
cx, cy, cw, ch, gap = Inches(4.1), Inches(4.35), Inches(3.9), Inches(1.25), Inches(0.3)
for i, (val, lbl) in enumerate(cards):
    col, row = i % 2, i // 2
    x = cx + col * (cw + gap)
    y = cy + row * (ch + Inches(0.25))
    stat_card(s, x, y, cw, ch, val, lbl)

s = new_content_slide(2, "The Proof", "Worked With, and Worth Trusting", title_size=32)
textbox(s, Inches(0.72), Inches(1.9), Inches(11.9), Inches(0.6),
        "Companies and tools worked with include HeyGen, Higgsfield, Wondershare, and\n"
        "Hostinger — plus guest masterclasses on major technology channels.",
        size=15, color=MUTED, line_spacing=1.3)
logos = ["Kashif Majeed", "HBA Services", "Lets Uncover", "Meet Mughals"]
lx, lw, lgap = Inches(0.72), Inches(2.75), Inches(0.3)
for i, name in enumerate(logos):
    x = lx + i * (lw + lgap)
    placeholder_box(s, x, Inches(2.9), lw, Inches(1.6), name.upper() + " LOGO")
rect(s, Inches(0.72), Inches(4.9), Inches(11.9), Inches(1.1), fill=WHITE, line_color=CARD_BORDER, radius=0.06, shadow=True)
textbox(s, Inches(1.05), Inches(5.15), Inches(11.2), Inches(0.6),
        "Credibility principle — show evidence rather than spending too much time describing yourself.",
        size=15, color=INK, bold=True)

s = new_content_slide(2, "Real Results", "In Their Own Words", title_size=32)
vids = ["Testimonial 1", "Testimonial 2", "Testimonial 3"]
vw, vgap = Inches(3.75), Inches(0.3)
for i, name in enumerate(vids):
    x = Inches(0.72) + i * (vw + vgap)
    placeholder_box(s, x, Inches(2.1), vw, Inches(4.2), name.upper() + " VIDEO")

# =============================================================================
# SECTION 3 — THE AGENDA
# =============================================================================

s = new_content_slide(3, "The Agenda", "Topics We Will Discuss", title_size=34)
roadmap = ["AI Basics & Mental Models", "Prompting with a Repeatable Framework",
           "Perplexity Comet — The Agentic Browser", "Custom GPTs — Build Your Own AI Assistant",
           "Claude — Build a Mini App Live", "OpenClaw — The AGI Principle",
           "HeyGen — AI Video Creation", "Automation with n8n — Agents vs Automation",
           "Selling AI Skills & Services"]
rmy = Inches(2.0)
half = -(-len(roadmap) // 2)
for i, item in enumerate(roadmap):
    col, row = i // half, i % half
    x = Inches(0.72) + col * Inches(6.1)
    y = rmy + Inches(0.66) * row
    textbox(s, x, y, Inches(0.6), Inches(0.5), f"{i + 1:02d}", size=18, color=ORANGE, bold=True, font=HEAD_FONT)
    textbox(s, x + Inches(0.65), y + Inches(0.02), Inches(5.3), Inches(0.5), item, size=15, color=INK, anchor=MSO_ANCHOR.TOP)
live_badge(s, on_dark=False)

# =============================================================================
# SECTION 4 — AI BASICS
# =============================================================================

s = section_break(BG_LIGHT, 4, "AI Basics", "Mental Models, Not Theory",
                   "The goal is not to teach AI theory in depth — it's to give beginners a mental\nmodel that makes the rest of the workshop understandable.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(4, "What is AI?", "AI Has Been Invisible in Your Life for Years", title_size=28)
textbox(s, Inches(0.72), Inches(1.85), Inches(11.9), Inches(0.9),
        "AI is a broad category of software that can recognize patterns, learn from\n"
        "data, and produce useful outputs — not one single tool or product.",
        size=15, color=MUTED, line_spacing=1.3)
ai_ex = ["A map app reroutes you around traffic", "A bank flags an unusual purchase",
         "A customer-support chatbot answers common questions", "A recommendation system predicts what you'll want next"]
chip_grid(s, ai_ex, Inches(0.72), Inches(2.9), Inches(11.9), Inches(2.6), cols=2, size=15)

s = new_content_slide(4, "The Mental Model", "AI → Model → LLM → ChatGPT", title_size=32)
chain = ["AI — the overall field / category", "Models — trained systems that perform particular tasks",
         "LLMs — models specialized in understanding and generating language",
         "ChatGPT — a product/interface that lets people use an LLM effectively"]
step_list(s, chain, Inches(0.72), Inches(2.0), Inches(11.5), row_h=0.85, size=17)
textbox(s, Inches(0.72), Inches(5.6), Inches(11.9), Inches(0.9),
        "Key takeaway — don't think of AI as one tool. Think of it as an ecosystem of\nmodels and products designed for different kinds of work.",
        size=14, color=MUTED, bold=True, line_spacing=1.3)

s = new_content_slide(4, "Evolution", "From Rules to Generative AI", title_size=30)
timeline = ["1990s — Spam filters and rule-based systems", "1997 — Deep Blue defeats chess champion Garry Kasparov",
            "2000s — Search ranking and information retrieval mature",
            "2010s — Recommendation systems, computer vision, Face ID, speech recognition mature",
            "2020s — Generative AI produces text, images, audio, video, and code"]
bullets(s, timeline, Inches(0.72), Inches(1.95), Inches(11.5), Inches(3.6), size=16, gap=0.22)
textbox(s, Inches(0.72), Inches(5.9), Inches(11.9), Inches(0.8),
        "Why now? Three pieces came together: enormous computing power, huge amounts\nof data, and modern architectures like transformers.",
        size=14, color=MUTED, bold=True, line_spacing=1.3)

s = new_content_slide(4, "The AI Hierarchy", "AI → ML → Deep Learning → GenAI / LLMs", title_size=28)
layers = [("AI", "The broad outer layer"), ("Machine Learning", "Learns patterns from data"),
          ("Deep Learning", "ML using neural networks"), ("Generative AI / LLMs", "The innermost layer, creates new content")]
lyw, lygap = Inches(2.85), Inches(0.2)
for i, (name, desc) in enumerate(layers):
    x = Inches(0.72) + i * (lyw + lygap)
    fill = ORANGE if i == len(layers) - 1 else WHITE
    txt = WHITE if i == len(layers) - 1 else INK
    rect(s, x, Inches(2.3), lyw, Inches(2.6), fill=fill, line_color=CARD_BORDER if fill == WHITE else None, radius=0.08, shadow=True)
    textbox(s, x + Inches(0.2), Inches(2.6), lyw - Inches(0.4), Inches(0.9), name, size=16, color=txt, bold=True, font=HEAD_FONT, line_spacing=1.05)
    textbox(s, x + Inches(0.2), Inches(3.7), lyw - Inches(0.4), Inches(1.1), desc, size=12, color=txt if fill == ORANGE else MUTED, line_spacing=1.25)
textbox(s, Inches(0.72), Inches(5.3), Inches(11.9), Inches(0.7),
        "Takeaway — these aren't separate competing categories, they're nested layers of one stack.",
        size=14, color=MUTED, bold=True)

s = new_content_slide(4, "Machine Learning & Deep Learning", "Teaching by Examples", title_size=27)
qa_row(s, Inches(0.72), Inches(1.95), Inches(11.9), "Traditional programming",
       "RULES → COMPUTER → RESULT", "You write every rule by hand, one at a time.", power=False)
qa_row(s, Inches(0.72), Inches(3.15), Inches(11.9), "Machine learning",
       "EXAMPLES → MODEL LEARNS PATTERNS → RESULT", "You show it data instead, and it finds the rules itself.", power=True)
textbox(s, Inches(0.72), Inches(4.5), Inches(11.9), Inches(0.6),
        "Deep learning is ML using neural networks: layer 1 detects edges, layer 2 shapes, layer 3\nfeatures — real-world examples include Face ID, translation, and image recognition.",
        size=14, color=MUTED, line_spacing=1.3)

s = new_content_slide(4, "LLMs & How They Work", "Large Language Model", title_size=30)
textbox(s, Inches(0.72), Inches(1.95), Inches(11.9), Inches(0.9),
        "An LLM is a type of AI model trained on very large amounts of text and other\n"
        "data to understand and generate language — the models behind ChatGPT and Claude.",
        size=15, color=MUTED, line_spacing=1.3)
rule_pairs = [("Mental model", "An LLM takes the context you give it, processes the relationships between\nthe tokens, and generates a response one token at a time."),
              ("Important distinction", "ChatGPT or Claude are applications that use models — an LLM is the\nunderlying model technology."),
              ("Core takeaway", "An LLM generates language by repeatedly predicting the next token based\non context — not simple “autocomplete.”")]
ry = Inches(3.1)
for i, (lbl, txt) in enumerate(rule_pairs):
    y = ry + Inches(1.25) * i
    rect(s, Inches(0.72), y, Inches(11.9), Inches(1.05), fill=WHITE, line_color=CARD_BORDER, radius=0.06, shadow=True)
    rich_textbox(s, Inches(1.0), y + Inches(0.18), Inches(11.3), Inches(0.75),
                 [(lbl + "  ", ORANGE, True), (txt, MUTED, False)], size=13, line_spacing=1.25)

s = new_content_slide(4, "Narrow AI vs. AGI, and Limitations", "What AI Can — and Can't — Do", title_size=28)
qa_row(s, Inches(0.72), Inches(1.95), Inches(11.9), "Narrow AI",
       "WHAT EXISTS TODAY", "Powerful at defined tasks, not generally intelligent across every domain.", power=True)
qa_row(s, Inches(0.72), Inches(3.15), Inches(11.9), "AGI",
       "NOT A SETTLED, EXISTING SYSTEM", "The long-term idea of broad, human-level intelligence that transfers across domains.", power=False)
lims = ["Hallucinations / incorrect information", "Knowledge may be outdated depending on the model/tool",
        "Prompt and context quality affect output quality", "AI can sound confident even when wrong"]
bullets(s, lims, Inches(0.72), Inches(4.5), Inches(11.9), Inches(2.0), size=15, gap=0.16)
textbox(s, Inches(0.72), Inches(6.6), Inches(11.9), Inches(0.5),
        "Rule — use AI for leverage, but verify important information.", size=13, color=MUTED, bold=True)

# =============================================================================
# SECTION 5 — PROMPTING
# =============================================================================

s = section_break(BG_LIGHT, 5, "Prompting", "Stop Guessing What to Type",
                   "A repeatable framework for talking to AI — one you'll reuse for the rest of today.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(5, "Core Principle", "Bad Input, Unpredictable Output", title_size=32)
textbox(s, Inches(0.72), Inches(2.2), Inches(11.0), Inches(1.0),
        "Better context + clearer instructions → better output.", size=24, color=INK, bold=True, font=HEAD_FONT)
textbox(s, Inches(0.72), Inches(3.3), Inches(11.0), Inches(1.2),
        "Prompting is increasingly part of a broader skill: giving AI the right task,\ncontext, examples, constraints, and output requirements.",
        size=17, color=MUTED, line_spacing=1.3)

s = new_content_slide(5, "The Demo Task · Level 1", "One Task, Three Prompt Levels", title_size=27)
code_block(s, Inches(0.72), Inches(1.85), Inches(11.9), Inches(1.1),
           ["Create a HeyGen script for a short promotional video for our", "AI workshop."], label="The Task")
code_block(s, Inches(0.72), Inches(3.15), Inches(5.75), Inches(1.0),
           ["Write a script for my AI workshop video."], label="Level 1 — Weak", title_color=RGBColor(0xFF, 0x8A, 0x80))
weak_pts = ["No clear objective", "No audience", "No tone or style", "No length / structure", "No CTA or success criteria"]
bullets(s, weak_pts, Inches(6.75), Inches(3.2), Inches(5.9), Inches(2.2), size=13, gap=0.1)

s = new_content_slide(5, "The Demo Task · Level 2", "Good Prompt", title_size=30)
code_block(s, Inches(0.72), Inches(1.9), Inches(11.9), Inches(2.1),
           ["Write a 60-second promotional script for an AI workshop. The",
            "audience is beginners who want to learn practical AI skills. Make",
            "it engaging, easy to understand, and persuasive. Explain what they",
            "will learn and end with a clear call to action to register."], label="Level 2 — Good", title_color=CODE_ORANGE)
good_pts = ["The task is clearer", "Audience is defined", "Length is specified", "Tone and objective are clearer", "CTA is included"]
chip_grid(s, good_pts, Inches(0.72), Inches(4.3), Inches(11.9), Inches(1.9), cols=3, size=13)

s = new_content_slide(5, "The Demo Task · Level 3", "The 4D Framework Prompt", title_size=30)
textbox(s, Inches(0.72), Inches(1.7), Inches(11.0), Inches(0.4), "Define → Describe → Demonstrate → Deliver", size=15, color=ORANGE, bold=True)
meta_lines = [
    "DEFINE  Create a 60-second promotional video script for our AI",
    "workshop. Motivate viewers to register.",
    "DESCRIBE  Audience: beginners, professionals, creators, freelancers,",
    "business owners overwhelmed by AI tools. Keep language beginner-",
    "friendly, energetic, credible, easy to speak on camera.",
    "DEMONSTRATE  Structure: 1) Hook 2) Audience's problem 3) What",
    "they'll learn 4) One concrete transformation 5) Strong CTA.",
    "Avoid generic AI buzzwords — write conversationally.",
    "DELIVER  One polished 60-second HeyGen-ready script. Short spoken",
    "sentences, natural pauses. End with a direct registration CTA.",
]
code_block(s, Inches(0.72), Inches(2.2), Inches(11.9), Inches(4.85), meta_lines, label="Level 3 — 4D")

s = new_content_slide(5, "Live Comparison", "All Three, Side by Side", title_size=30)
qa_row(s, Inches(0.72), Inches(2.0), Inches(11.9), "Bad prompt", "GENERIC AND UNPREDICTABLE",
       "No task, no audience, no constraints — the model has to guess.", power=False)
qa_row(s, Inches(0.72), Inches(3.2), Inches(11.9), "Good prompt", "MUCH MORE USEFUL",
       "Task and context are clearer, but still generic in places.", power=False)
qa_row(s, Inches(0.72), Inches(4.4), Inches(11.9), "4D prompt", "STRUCTURED AND SPECIFIC",
       "Repeatable, and aligned to the exact output needed.", power=True)
textbox(s, Inches(0.72), Inches(5.7), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)

s = new_content_slide(5, "The Loop · Go Deeper", "The 4D Prompt Becomes Our Workshop Tool", title_size=27)
loop_steps = ["Learn the 4D framework", "Build the prompt", "Compare results",
              "Select the strongest version", "Use it to create the HeyGen video in Section 10"]
step_list(s, loop_steps, Inches(0.72), Inches(1.95), Inches(11.5), row_h=0.62, size=14, circle_d=0.4)
textbox(s, Inches(0.72), Inches(5.4), Inches(11.9), Inches(0.5),
        "The point — prompting isn't theory, it directly improves the work they produce.",
        size=13, color=MUTED, bold=True)
links = ["Anthropic · Learn / Build with Claude", "Anthropic · AI Fluency Framework", "Bonus · Workshop prompt library"]
chip_grid(s, links, Inches(0.72), Inches(6.0), Inches(11.9), Inches(0.9), cols=3, size=12, fill=RGBColor(0xF0, 0xE8, 0xDC), border=ORANGE)

# =============================================================================
# SECTION 6 — PERPLEXITY COMET
# =============================================================================

s = section_break(BG_LIGHT, 6, "Perplexity Comet", "The Agentic Browser",
                   "From AI that answers questions to AI that can perform multi-step browser tasks.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(6, "Search vs. Delegate", "What Is an Agentic Browser?", title_size=30)
qa_row(s, Inches(0.72), Inches(2.1), Inches(11.9), "Traditional search", "YOU DO THE WORK",
       "You search → read → click → compare → decide, one step at a time.", power=False, h=1.3)
qa_row(s, Inches(0.72), Inches(3.7), Inches(11.9), "Agentic browser", "THE AGENT DOES THE WORK",
       "You give a goal — the agent navigates, researches, compares, and completes\nsteps with less manual work.", power=True, h=1.3)

s = new_content_slide(6, "What This Looks Like", "Comet, Mid-Task", title_size=32)
placeholder_box(s, Inches(0.72), Inches(2.0), Inches(11.9), Inches(4.2), "COMET DEMO SCREENSHOT")

s = new_content_slide(6, "Live Demo", "Delegate a Real Task", title_size=32)
tasks = ["Research competitors", "Compare products", "Find potential leads", "Research a travel plan", "Gather info from multiple sites and summarize it"]
chip_grid(s, tasks, Inches(0.72), Inches(2.0), Inches(11.9), Inches(2.6), cols=3, size=14)
textbox(s, Inches(0.72), Inches(5.0), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)
textbox(s, Inches(0.72), Inches(5.55), Inches(11.9), Inches(0.6),
        "Teaching point — the value is delegating a multi-step knowledge task, not the browser itself.",
        size=14, color=MUTED, bold=True)

# =============================================================================
# SECTION 7 — CUSTOM GPTs
# =============================================================================

s = section_break(BG_LIGHT, 7, "Custom GPTs", "Build Your Own AI Assistant",
                   "Stop repeating the same instructions every day.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(7, "Why Custom Assistants?", "A Specialized Assistant Has", title_size=32)
why_gpt = ["A clear role", "Persistent instructions", "Reference files / knowledge", "A defined workflow", "Output rules"]
bullets(s, why_gpt, Inches(0.72), Inches(2.0), Inches(11.5), Inches(3.0), size=18, gap=0.25)

s = new_content_slide(7, "The Builder", "Where You'll Set All This Up", title_size=30)
placeholder_box(s, Inches(0.72), Inches(1.95), Inches(5.6), Inches(3.6), "CUSTOM GPT BUILDER UI")
builder_pts = ["One place to name the assistant and define its job", "Instructions live here permanently — no re-typing per chat",
               "Attach knowledge files it should always have access to", "Test it right inside the builder before sharing it"]
bullets(s, builder_pts, Inches(6.65), Inches(2.1), Inches(5.9), Inches(3.4), size=15, gap=0.25)
textbox(s, Inches(0.72), Inches(5.75), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)

s = new_content_slide(7, "Live Build", 'Build "AI Content Strategist"', title_size=30)
gpt_steps = ["Define the job", "Write the instructions", "Add knowledge / files", "Add examples", "Test with real inputs", "Improve based on failures"]
step_list(s, gpt_steps, Inches(0.72), Inches(1.95), Inches(11.5), row_h=0.58, size=15, circle_d=0.42)
textbox(s, Inches(0.72), Inches(5.7), Inches(11.9), Inches(0.7),
        "Key idea — a custom GPT isn't just a fancy prompt, it's a reusable interface around a repeatable task.",
        size=14, color=MUTED, bold=True)

# =============================================================================
# SECTION 8 — CLAUDE: BUILD A MINI APP LIVE
# =============================================================================

s = section_break(DARK_BG, 8, "Claude · The Hero Demo", "Build a Mini App, Live",
                   "Slow down here. Let the room see the transformation.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(8, "Why Claude for Building", "What Makes Claude Useful for Building?", title_size=27)
claude_pts = ["Writing and reasoning", "Working with long context", "Understanding project requirements",
              "Generating and modifying code", "Iterating on a working product"]
bullets(s, claude_pts, Inches(0.72), Inches(2.0), Inches(11.5), Inches(3.2), size=18, gap=0.25)

s = new_content_slide(8, "Build Process", "From Idea to Working App", title_size=30)
code_block(s, Inches(0.72), Inches(1.85), Inches(11.9), Inches(1.0),
           ["I want a simple app for [ask the room what to build]."], label="Starting Idea")
build_steps = ["Ask the room for a real business or idea, live", "Ask Claude to clarify audience, goals, features, journey",
               "Generate the first version", "Run the app / preview the result",
               "Give iterative feedback", "Add or change features live", "Polish the UI", "Show the final working result"]
step_list(s, build_steps[:4], Inches(0.72), Inches(3.15), Inches(5.7), row_h=0.62, size=14, circle_d=0.4)
step_list(s, build_steps[4:], Inches(6.65), Inches(3.15), Inches(5.7), row_h=0.62, size=14, circle_d=0.4, start=5)
textbox(s, Inches(0.72), Inches(6.15), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)

s = new_content_slide(8, "Teaching Moment", "Vibe Coding", title_size=32)
placeholder_box(s, Inches(0.72), Inches(1.95), Inches(5.3), Inches(2.6), "VIBE CODING CYCLE DIAGRAM")
vibe_pts = ["Idea → Describe → Generate → Run", "Inspect → Give feedback → Iterate", "Repeat until it's right"]
bullets(s, vibe_pts, Inches(6.4), Inches(2.1), Inches(6.2), Inches(2.4), size=17, gap=0.3)
textbox(s, Inches(0.72), Inches(4.9), Inches(11.9), Inches(1.0),
        "Still required — you don't need to manually write every line of code, but you\nstill need to understand the product, requirements, testing, and quality.",
        size=15, color=MUTED, line_spacing=1.3)

s = new_content_slide(8, "The Bigger Lesson · Bonus", "It's Not About the App You Just Watched", title_size=27)
textbox(s, Inches(0.72), Inches(1.95), Inches(11.9), Inches(1.0),
        "The lesson — if you can clearly describe a problem, you can increasingly use\nAI to help turn that problem into a working prototype.",
        size=18, color=INK, bold=True, line_spacing=1.3)
textbox(s, Inches(0.72), Inches(3.15), Inches(11.0), Inches(0.4), "THE REST OF THE CLAUDE TOOLKIT", size=13, color=ORANGE, bold=True, spacing=1.5)
toolkit = [("Projects", "A dedicated workspace that holds your files, instructions, and context across every chat"),
           ("Artifacts", "Docs, code, and mini-apps Claude builds in a live side panel — the app you just watched, lives here"),
           ("Claude Code", "Claude working directly in your terminal and codebase — reads files, runs commands, tests its own changes")]
tkw, tkgap = Inches(3.7), Inches(0.35)
for i, (name, desc) in enumerate(toolkit):
    x = Inches(0.72) + i * (tkw + tkgap)
    win = i == 1
    rect(s, x, Inches(3.65), tkw, Inches(3.0), fill=RGBColor(0xEE, 0xF8, 0xE4) if win else WHITE,
         line_color=GREEN if win else CARD_BORDER, line_w=Pt(1.25), radius=0.07, shadow=True)
    textbox(s, x + Inches(0.25), Inches(3.95), tkw - Inches(0.5), Inches(0.5), name, size=18, color=INK, bold=True, font=HEAD_FONT)
    textbox(s, x + Inches(0.25), Inches(4.55), tkw - Inches(0.5), Inches(1.9), desc, size=13, color=MUTED, line_spacing=1.3)

# =============================================================================
# SECTION 9 — OPENCLAW
# =============================================================================

s = section_break(DARK_BG, 9, "OpenClaw", "The AGI Principle", None)
oc_steps = ["AI answers my question", "AI uses tools to complete a task", "AI can operate toward a goal with less step-by-step supervision"]
step_list(s, oc_steps, Inches(0.72), Inches(5.05), Inches(11.5), row_h=0.55, size=14, circle_d=0.4, color=BG_LIGHT)
placeholder_box(s, Inches(9.6), Inches(1.1), Inches(2.9), Inches(2.9), "OPENCLAW MASCOT")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(9, "Key Point", "An Agent Is More Than a Chatbot", title_size=30)
agent_pts = ["It can be given a goal", "It can access and use tools", "It decides what steps to take",
             "It observes its own results", "It continues until the task is complete", "All within defined permissions and constraints"]
chip_grid(s, agent_pts, Inches(0.72), Inches(2.0), Inches(11.9), Inches(2.6), cols=3, size=14)
textbox(s, Inches(0.72), Inches(5.0), Inches(11.9), Inches(0.7),
        "Important — AGI and agentic AI are not the same thing. An agent can be highly\nautonomous without being AGI.",
        size=15, color=MUTED, bold=True, line_spacing=1.3)

# =============================================================================
# SECTION 10 — HEYGEN
# =============================================================================

s = section_break(BG_LIGHT, 10, "HeyGen", "AI Video Creation",
                   "AI avatars, cloned voices, and instant localization.")
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(10, "The Capabilities", "What Can AI Video Do?", title_size=32)
caps = ["AI avatars", "Voice generation / cloning", "Multilingual localization", "Training and educational videos",
        "Sales videos", "Marketing content", "Scalable content production"]
chip_grid(s, caps, Inches(0.72), Inches(2.0), Inches(11.9), Inches(3.6), cols=4, size=14)

s = new_content_slide(10, "The Script", "From Section 5's 4D Prompt", title_size=28)
hg_lines = [
    "Stop scrolling past AI. In the next two hours, you're going to use it,",
    "build with it, and learn how to actually make money with it.",
    "",
    "This workshop fixes the noise. You'll learn the mental models that",
    "matter, a prompting framework you'll reuse forever, and how to build",
    "a real working app with Claude — live, from a single sentence.",
    "",
    "Then we go further — AI video, real automation with n8n, and how",
    "to turn everything you learn into paid work.",
    "",
    "Two hours. Nine skills. One session. Register now.",
]
code_block(s, Inches(0.72), Inches(1.85), Inches(11.9), Inches(4.9), hg_lines, label="HeyGen Script")

s = new_content_slide(10, "Live Demo · Inside the Studio", "Script → Finished Video", title_size=28)
placeholder_box(s, Inches(0.72), Inches(1.95), Inches(5.6), Inches(3.6), "HEYGEN STUDIO INTERFACE")
studio_pts = ["My Dashboard — every project, avatar, and render in one place",
              "My Avatar / My Voice — the two libraries picked live",
              "Create Video — where the script from Section 5 goes in",
              "Ready-made templates for ads, promos, and explainers"]
bullets(s, studio_pts, Inches(6.65), Inches(2.1), Inches(5.9), Inches(3.4), size=14, gap=0.25)
textbox(s, Inches(0.72), Inches(5.75), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)

s = new_content_slide(10, "The Result", "From Script to Finished Video", title_size=30)
placeholder_box(s, Inches(0.72), Inches(1.95), Inches(11.9), Inches(3.3), "FINISHED HEYGEN VIDEO")
textbox(s, Inches(0.72), Inches(5.5), Inches(11.9), Inches(1.0),
        "Same script from Section 5's 4D prompt. Same avatar and voice picked live.\nRendered in minutes — ready to post, no camera or crew involved.",
        size=15, color=MUTED, line_spacing=1.3)
textbox(s, Inches(0.72), Inches(6.6), Inches(11.9), Inches(0.5),
        "The payoff — this is what “one script, many videos” actually looks like.", size=15, color=INK, bold=True)

s = new_content_slide(10, "The Business Angle", "One Script, Many Outcomes", title_size=30)
outcomes = ["One script → multiple languages", "One idea → many videos", "Faster content production",
            "Lower production overhead", "Can become part of an automated content system"]
chip_grid(s, outcomes, Inches(0.72), Inches(2.0), Inches(11.9), Inches(2.6), cols=3, size=14)
rect(s, Inches(0.72), Inches(5.0), Inches(11.9), Inches(1.0), fill=RGBColor(0x24, 0x20, 0x1C), radius=0.08, shadow=True)
textbox(s, Inches(0.72), Inches(5.28), Inches(11.9), Inches(0.5),
        "Bridge — but what if we don't want to do these steps manually every time?",
        size=17, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

# =============================================================================
# SECTION 11 — AUTOMATION WITH N8N
# =============================================================================

s = section_break(BG_LIGHT, 11, "Automation With n8n", "Agents vs Automation", None)
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(11, "What Is Automation?", "A Defined Process", title_size=32)
auto_steps = ["Trigger", "Step 1", "Step 2", "Step 3", "Result"]
step_list(s, auto_steps, Inches(0.72), Inches(2.0), Inches(6.0), row_h=0.7, size=16, circle_d=0.48)
textbox(s, Inches(0.72), Inches(5.9), Inches(11.9), Inches(0.7),
        "Example — new form submission → add lead to CRM → send email → notify sales team.",
        size=14, color=MUTED, bold=True)

s = new_content_slide(11, "What Is an AI Agent?", "A Flexible, Goal-Driven Process", title_size=30)
placeholder_box(s, Inches(0.72), Inches(1.95), Inches(5.3), Inches(2.8), "AGENT NAVIGATING TOWARD GOAL")
agent_flow = ["Goal → decide steps → use tools", "Observe results → adapt → complete task"]
bullets(s, agent_flow, Inches(6.4), Inches(2.3), Inches(6.2), Inches(2.0), size=17, gap=0.3)

s = new_content_slide(11, "Side by Side", "Automation vs. Agent", title_size=32)
table_block(s, Inches(0.72), Inches(2.0), Inches(11.9), Inches(3.0), ["Automation", "AI Agent"],
            [["Predefined workflow", "Goal-driven workflow"],
             ["Predictable steps", "Can choose next steps"],
             ["Great for repetitive processes", "Great for complex/variable tasks"],
             ["Easier to test", "More flexible, but needs stronger controls"]])

s = new_content_slide(11, "Live Demo", "One Real n8n Workflow", title_size=32)
textbox(s, Inches(0.72), Inches(1.95), Inches(11.5), Inches(0.7),
        "Lead comes in → AI analyzes lead → classifies intent → generates personalized\nresponse → stores lead → notifies salesperson.",
        size=14, color=MUTED, line_spacing=1.3)
n8n_steps = ["Trigger", "Data transformation", "AI decision / generation step", "Action", "Final result"]
step_list(s, n8n_steps, Inches(0.72), Inches(3.15), Inches(6.0), row_h=0.65, size=15, circle_d=0.45)
placeholder_box(s, Inches(7.15), Inches(3.1), Inches(5.4), Inches(3.3), "N8N WORKFLOW CANVAS")
textbox(s, Inches(0.72), Inches(6.55), Inches(4), Inches(0.4), "LIVE DEMO", size=13, color=ORANGE, bold=True, spacing=1.5)

s = new_content_slide(11, "The Bigger Mental Model", "Four Building Blocks", title_size=32)
blocks = [("Prompt", "One task"), ("Custom GPT", "Reusable assistant"),
          ("Agent", "AI that pursues a goal using tools"), ("Automation", "Repeatable system that runs without manual effort")]
bw4, bgap4 = Inches(2.85), Inches(0.2)
for i, (name, desc) in enumerate(blocks):
    x = Inches(0.72) + i * (bw4 + bgap4)
    win = i == 3
    rect(s, x, Inches(2.1), bw4, Inches(2.3), fill=RGBColor(0xEE, 0xF8, 0xE4) if win else WHITE,
         line_color=GREEN if win else CARD_BORDER, line_w=Pt(1.25), radius=0.08, shadow=True)
    textbox(s, x + Inches(0.2), Inches(2.4), bw4 - Inches(0.4), Inches(0.5), name, size=17, color=INK, bold=True, font=HEAD_FONT)
    textbox(s, x + Inches(0.2), Inches(3.0), bw4 - Inches(0.4), Inches(1.2), desc, size=12, color=MUTED, line_spacing=1.25)
textbox(s, Inches(0.72), Inches(4.8), Inches(11.9), Inches(0.6),
        "Combine them — AI-powered automation = agent + automation, together.", size=15, color=INK, bold=True)

# =============================================================================
# SECTION 12 — SELLING AI SKILLS
# =============================================================================

s = section_break(BG_LIGHT, 12, "Selling", "Turn AI Skills Into Money",
                   '"Okay, how do I actually make money with this?"')
page_num[0] += 1; footer_page(s, page_num[0])

s = new_content_slide(12, "What Are People Buying?", 'Clients Don\'t Buy "AI" — They Buy Outcomes', title_size=27)
buys = ["More leads", "More content", "Faster research", "Lower operating costs",
        "Better customer support", "Automated workflows", "Internal tools", "Faster product development"]
chip_grid(s, buys, Inches(0.72), Inches(1.95), Inches(11.9), Inches(3.6), cols=4, size=13)

s = new_content_slide(12, "Beginner-Friendly Offers", "Offers You Can Start With", title_size=30)
offers = ["AI content systems", "AI video creation", "Custom AI assistants", "Research systems",
          "Lead-generation workflows", "n8n automations", "AI-powered internal tools", "AI website/app prototypes"]
chip_grid(s, offers, Inches(0.72), Inches(1.95), Inches(11.9), Inches(3.6), cols=4, size=13)

s = new_content_slide(12, "The Offer Formula", "Skill → Problem → Outcome → Offer → Proof", title_size=27)
qa_row(s, Inches(0.72), Inches(2.1), Inches(11.9), '"I know n8n and ChatGPT."', "WEAK",
       "A skill list. Says nothing about the outcome a client gets.", power=False, h=1.2)
qa_row(s, Inches(0.72), Inches(3.55), Inches(11.9), '"I build an AI lead-follow-up system\nthat responds within minutes."', "STRONG",
       "A named outcome, for a specific problem, that a client can picture immediately.", power=True, h=1.5)

s = new_content_slide(12, "First Client", "Getting Your First Client", title_size=30)
client_steps = ["Pick one niche", "Pick one painful problem", "Build one demo", "Package it as a simple offer",
                "Contact potential clients", "Show the demo instead of only talking about AI",
                "Start with a small implementation", "Turn the result into a case study"]
step_list(s, client_steps[:4], Inches(0.72), Inches(1.95), Inches(5.7), row_h=0.6, size=13, circle_d=0.4)
step_list(s, client_steps[4:], Inches(6.65), Inches(1.95), Inches(5.7), row_h=0.6, size=13, circle_d=0.4, start=5)

s = new_content_slide(12, "Final Challenge", "Pick One Problem You Can Solve This Week", title_size=27)
textbox(s, Inches(0.72), Inches(2.3), Inches(11.5), Inches(2.0),
        "Final message — you don't need to master every AI tool. You need to\n"
        "understand the fundamentals, learn how to communicate with AI, choose the\n"
        "right tools, build useful systems, and connect those systems to real problems.",
        size=19, color=INK, line_spacing=1.4)

# =============================================================================
# SECTION 13 — CLOSING
# =============================================================================

s = new_content_slide(13, "Closing", "You Now Know How to Use, Build,\nAutomate, and Sell With AI", title_size=28)
recap13 = ["AI fundamentals & mental models", "A repeatable prompting framework", "An agentic browser doing real research",
           "Your own custom AI assistant", "A working app, built live with Claude", "The difference between an agent and automation",
           "AI video creation with HeyGen", "A real n8n automation workflow", "How to turn all of it into paid work"]
chip_grid(s, recap13, Inches(0.72), Inches(2.7), Inches(11.9), Inches(3.6), cols=3, size=13)

s = add_slide(bg=DARK_BG)
kicker(s, "Closing", color=ORANGE)
progress_tag(s, 13, on_dark=True)
headline(s, "Live Q&A", y=Inches(2.8), size=56, color=BG_LIGHT)
underline(s, Inches(0.72), Inches(4.1), Inches(2.2))
live_badge(s, on_dark=True)
page_num[0] += 1; footer_page(s, page_num[0])

# ------------------------------------------------------------------- save --
import os
out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "AI-Mastery-Workshop.pptx")
prs.save(out_path)
print(f"Saved {len(prs.slides)} slides -> {out_path}")
